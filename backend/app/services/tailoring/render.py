"""On-demand resume rendering: single-column ATS-friendly DOCX and PDF (owner: tailoring-agent).

Never stores generated files on disk. Renders dynamically from TailoredResume JSON.
ATS rules: Single column, clear semantic headers, no tables, no icons, standard margins.
"""

from __future__ import annotations

import io

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from xhtml2pdf import pisa

from app.schemas import ParsedResume, TailoredResume


def _get_contact_line(profile: ParsedResume | None) -> str:
    parts: list[str] = []
    if profile:
        if profile.email:
            parts.append(profile.email)
        if profile.phone:
            parts.append(profile.phone)
    return " | ".join(parts)


def render_docx(tailored: TailoredResume, profile: ParsedResume | None = None) -> bytes:
    """Render a single-column, ATS-friendly DOCX file from tailored resume content."""
    doc = Document()

    # 0.8-inch margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Name Header
    name = (profile.name if profile and profile.name else "Candidate").strip()
    p_name = doc.add_paragraph()
    r_name = p_name.add_run(name)
    r_name.font.name = "Calibri"
    r_name.font.size = Pt(20)
    r_name.font.bold = True
    p_name.paragraph_format.space_after = Pt(2)

    # Contact line
    contact = _get_contact_line(profile)
    if contact:
        p_contact = doc.add_paragraph()
        r_contact = p_contact.add_run(contact)
        r_contact.font.name = "Calibri"
        r_contact.font.size = Pt(9.5)
        p_contact.paragraph_format.space_after = Pt(6)

    # Headline
    if tailored.headline:
        p_head = doc.add_paragraph()
        r_head = p_head.add_run(tailored.headline)
        r_head.font.name = "Calibri"
        r_head.font.size = Pt(11)
        r_head.font.bold = True
        p_head.paragraph_format.space_after = Pt(8)

    # Helper for adding standard section headings
    def add_section_header(title: str) -> None:
        p = doc.add_paragraph()
        r = p.add_run(title.upper())
        r.font.name = "Calibri"
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 41, 59)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)

    # 1. Professional Summary
    if tailored.summary:
        add_section_header("Professional Summary")
        p_sum = doc.add_paragraph()
        r_sum = p_sum.add_run(tailored.summary)
        r_sum.font.name = "Calibri"
        r_sum.font.size = Pt(10)
        p_sum.paragraph_format.space_after = Pt(6)

    # 2. Skills
    if tailored.skills_grouped:
        add_section_header("Technical Skills")
        for group_name, skills in tailored.skills_grouped.items():
            if skills:
                p_sk = doc.add_paragraph()
                r_lbl = p_sk.add_run(f"{group_name}: ")
                r_lbl.font.name = "Calibri"
                r_lbl.font.size = Pt(9.5)
                r_lbl.font.bold = True

                r_val = p_sk.add_run(", ".join(skills))
                r_val.font.name = "Calibri"
                r_val.font.size = Pt(9.5)
                p_sk.paragraph_format.space_after = Pt(2)

    # 3. Experience
    if tailored.experience:
        add_section_header("Work Experience")
        for exp in tailored.experience:
            p_exp = doc.add_paragraph()
            r_title = p_exp.add_run(exp.title)
            r_title.font.name = "Calibri"
            r_title.font.size = Pt(10.5)
            r_title.font.bold = True

            comp_str = f" - {exp.company}" if exp.company else ""
            dur_str = f" ({exp.duration})" if exp.duration else ""
            r_meta = p_exp.add_run(f"{comp_str}{dur_str}")
            r_meta.font.name = "Calibri"
            r_meta.font.size = Pt(10)
            p_exp.paragraph_format.space_after = Pt(2)

            for bullet in exp.bullets:
                p_b = doc.add_paragraph(style="List Bullet")
                r_b = p_b.add_run(bullet)
                r_b.font.name = "Calibri"
                r_b.font.size = Pt(9.5)
                p_b.paragraph_format.space_after = Pt(1.5)

    # 4. Projects
    if tailored.projects:
        add_section_header("Projects")
        for proj in tailored.projects:
            p_proj = doc.add_paragraph()
            r_pname = p_proj.add_run(proj.name)
            r_pname.font.name = "Calibri"
            r_pname.font.size = Pt(10.5)
            r_pname.font.bold = True
            p_proj.paragraph_format.space_after = Pt(2)

            if proj.description:
                p_desc = doc.add_paragraph()
                r_desc = p_desc.add_run(proj.description)
                r_desc.font.name = "Calibri"
                r_desc.font.size = Pt(9.5)
                p_desc.paragraph_format.space_after = Pt(2)

            for h in proj.highlights:
                p_h = doc.add_paragraph(style="List Bullet")
                r_h = p_h.add_run(h)
                r_h.font.name = "Calibri"
                r_h.font.size = Pt(9.5)
                p_h.paragraph_format.space_after = Pt(1.5)

    # 5. Education
    if tailored.education:
        add_section_header("Education")
        for edu in tailored.education:
            p_edu = doc.add_paragraph()
            r_inst = p_edu.add_run(edu.institution)
            r_inst.font.name = "Calibri"
            r_inst.font.size = Pt(10)
            r_inst.font.bold = True

            deg_parts = []
            if edu.degree:
                deg_parts.append(edu.degree)
            if edu.year:
                deg_parts.append(str(edu.year))
            if deg_parts:
                r_deg = p_edu.add_run(f" - {', '.join(deg_parts)}")
                r_deg.font.name = "Calibri"
                r_deg.font.size = Pt(9.5)
            p_edu.paragraph_format.space_after = Pt(2)

    # 6. Certifications
    if tailored.certifications:
        add_section_header("Certifications")
        for cert in tailored.certifications:
            p_cert = doc.add_paragraph(style="List Bullet")
            r_cname = p_cert.add_run(cert.name)
            r_cname.font.name = "Calibri"
            r_cname.font.size = Pt(9.5)
            if cert.issuer:
                r_iss = p_cert.add_run(f" ({cert.issuer})")
                r_iss.font.name = "Calibri"
                r_iss.font.size = Pt(9)
            p_cert.paragraph_format.space_after = Pt(1.5)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def render_pdf(tailored: TailoredResume, profile: ParsedResume | None = None) -> bytes:
    """Render a single-column, ATS-friendly PDF from tailored resume content."""
    name = (profile.name if profile and profile.name else "Candidate").strip()
    contact = _get_contact_line(profile)

    html_parts: list[str] = [
        "<!DOCTYPE html>",
        "<html><head><meta charset='utf-8'>",
        "<style>",
        "@page { size: A4; margin: 1.5cm; }",
        "body { font-family: Helvetica, Arial, sans-serif; font-size: 10pt; "
        "line-height: 1.35; color: #1e293b; }",
        "h1 { font-size: 20pt; margin: 0 0 4px 0; color: #0f172a; "
        "text-transform: uppercase; letter-spacing: 0.5px; }",
        ".contact { font-size: 9pt; color: #475569; margin-bottom: 6px; }",
        ".headline { font-size: 11pt; font-weight: bold; color: #0f766e; margin-bottom: 12px; }",
        "h2 { font-size: 11pt; border-bottom: 1.5px solid #0f172a; padding-bottom: 2px; "
        "margin: 12px 0 6px 0; text-transform: uppercase; letter-spacing: 0.5px; color: #0f172a; }",
        "p { margin: 0 0 6px 0; }",
        "ul { margin: 2px 0 6px 18px; padding: 0; }",
        "li { margin-bottom: 3px; font-size: 9.5pt; }",
        ".item-header { font-weight: bold; font-size: 10.5pt; color: #0f172a; }",
        ".item-sub { font-size: 9.5pt; color: #475569; margin-bottom: 3px; }",
        ".skills-group { margin-bottom: 4px; font-size: 9.5pt; }",
        ".group-title { font-weight: bold; color: #0f172a; }",
        "</style></head><body>",
        f"<h1>{name}</h1>",
    ]

    if contact:
        html_parts.append(f"<div class='contact'>{contact}</div>")
    if tailored.headline:
        html_parts.append(f"<div class='headline'>{tailored.headline}</div>")

    # Summary
    if tailored.summary:
        html_parts.append("<h2>Professional Summary</h2>")
        html_parts.append(f"<p>{tailored.summary}</p>")

    # Skills
    if tailored.skills_grouped:
        html_parts.append("<h2>Technical Skills</h2>")
        for group_name, skills in tailored.skills_grouped.items():
            if skills:
                skills_str = ", ".join(skills)
                html_parts.append(
                    f"<div class='skills-group'><span class='group-title'>"
                    f"{group_name}:</span> {skills_str}</div>"
                )

    # Experience
    if tailored.experience:
        html_parts.append("<h2>Work Experience</h2>")
        for exp in tailored.experience:
            comp_str = f" &bull; {exp.company}" if exp.company else ""
            dur_str = f" ({exp.duration})" if exp.duration else ""
            html_parts.append(f"<div class='item-header'>{exp.title}{comp_str}{dur_str}</div>")
            if exp.bullets:
                html_parts.append("<ul>")
                for b in exp.bullets:
                    html_parts.append(f"<li>{b}</li>")
                html_parts.append("</ul>")

    # Projects
    if tailored.projects:
        html_parts.append("<h2>Projects</h2>")
        for proj in tailored.projects:
            html_parts.append(f"<div class='item-header'>{proj.name}</div>")
            if proj.description:
                html_parts.append(f"<div class='item-sub'>{proj.description}</div>")
            if proj.highlights:
                html_parts.append("<ul>")
                for h in proj.highlights:
                    html_parts.append(f"<li>{h}</li>")
                html_parts.append("</ul>")

    # Education
    if tailored.education:
        html_parts.append("<h2>Education</h2>")
        for edu in tailored.education:
            deg_parts = []
            if edu.degree:
                deg_parts.append(edu.degree)
            if edu.year:
                deg_parts.append(str(edu.year))
            deg_str = f" &bull; {', '.join(deg_parts)}" if deg_parts else ""
            html_parts.append(f"<div class='item-header'>{edu.institution}{deg_str}</div>")

    # Certifications
    if tailored.certifications:
        html_parts.append("<h2>Certifications</h2><ul>")
        for cert in tailored.certifications:
            iss_str = f" ({cert.issuer})" if cert.issuer else ""
            html_parts.append(f"<li>{cert.name}{iss_str}</li>")
        html_parts.append("</ul>")

    html_parts.append("</body></html>")
    full_html = "\n".join(html_parts)

    pdf_buffer = io.BytesIO()
    pisa_status = pisa.CreatePDF(full_html, dest=pdf_buffer)
    if pisa_status.err:
        raise RuntimeError("xhtml2pdf rendering failed to generate PDF document.")

    return pdf_buffer.getvalue()
