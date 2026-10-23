"""Resume parser (owner: lead).

Baseline: text extraction + regex + taxonomy matching.
Enhanced (this PR): `.docx` file support (python-docx) and LLM-powered structured
extraction of education / experience / projects / summary via `app.core.llm`.

The regex/taxonomy baseline remains the deterministic fallback: any LLM failure or
schema-invalid LLM output falls back to it, so endpoints keep their normal responses.
"""

import io
import re

from docx import Document
from pypdf import PdfReader

from app.core.llm import complete_json
from app.core.taxonomy import extract_skills
from app.schemas import ParsedResume

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:\+91[\s-]?)?[6-9]\d{9}")

# Cap the text sent to the LLM so oversized resumes cannot blow up the prompt.
_MAX_LLM_TEXT_CHARS = 12000

_EXTRACT_SYSTEM = (
    "You extract structured data from resumes. Use only information present in the "
    "resume text. Never fabricate institutions, employers, projects, dates, degrees "
    "or any other qualification. If something is missing, omit it. Reply with JSON only."
)

_EXTRACT_PROMPT = """Extract structured information from the resume text below.
Return ONLY a JSON object with exactly this shape
(omit a field entirely if the resume does not contain it):
{{
  "summary": string | null,
  "education": [{{"institution": string, "degree": string | null, "year": string | null}}],
  "experience": [
    {{"title": string, "company": string | null, "summary": string | null,
      "duration": string | null}}
  ],
  "projects": [string]
}}
Rules:
- Use only information present in the resume. Never invent anything.
- Do not include "name", "email", "phone", "skills" or "raw_text"; they are handled separately.
- Every education entry needs an "institution" string;
  every experience entry needs a "title" string.

Resume text:
{text}
"""

# Only these LLM-provided fields are applied on top of the baseline result.
_LLM_FIELDS = ("summary", "education", "experience", "projects")


def extract_text(filename: str, data: bytes) -> str:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if lower.endswith(".docx"):
        return _extract_docx_text(data)
    return data.decode("utf-8", errors="ignore")


def _extract_docx_text(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    table_text: list[str] = []
    for table in document.tables:
        for row in table.rows:
            row_text = "  ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                table_text.append(row_text)
    return "\n".join(paragraphs + table_text)


def parse_resume(filename: str, data: bytes) -> ParsedResume:
    text = extract_text(filename, data).strip()
    baseline = _baseline_parse(text)
    try:
        llm_fields = _llm_extract(text)
    except Exception:
        return baseline
    return _merge(baseline, llm_fields)


def _baseline_parse(text: str) -> ParsedResume:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    email = EMAIL_RE.search(text)
    phone = PHONE_RE.search(text)
    return ParsedResume(
        name=lines[0] if lines else None,
        email=email.group(0) if email else None,
        phone=phone.group(0) if phone else None,
        skills=extract_skills(text),
        raw_text=text,
    )


def _llm_extract(text: str) -> dict[str, object]:
    """Attempt LLM extraction; validate the four supported fields against the schema.

    Raises (so callers can fall back) when the LLM is unavailable, returns invalid
    JSON, or the result does not validate against `ParsedResume`'s nested models.
    """
    if not text:
        raise ValueError("Refusing LLM extraction on an empty resume.")
    payload = complete_json(
        _EXTRACT_PROMPT.format(text=text[:_MAX_LLM_TEXT_CHARS]), _EXTRACT_SYSTEM
    )
    candidate = {
        key: payload[key] for key in _LLM_FIELDS if key in payload and payload[key] is not None
    }
    validated = ParsedResume(**candidate)
    return {key: getattr(validated, key) for key in _LLM_FIELDS}


def _merge(baseline: ParsedResume, llm_fields: dict[str, object]) -> ParsedResume:
    updates = {key: value or getattr(baseline, key) for key, value in llm_fields.items()}
    return baseline.model_copy(update=updates)
